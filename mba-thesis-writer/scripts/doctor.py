#!/usr/bin/env python3
"""Check a relocated skill without network, installation, or writes to the skill folder."""
import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]


def inspect(self_test=False):
    report = {"skill_root": str(ROOT), "python": sys.version.split()[0],
              "python_supported": sys.version_info >= (3, 10), "checks": [],
              "optional": {}, "errors": [], "visual_review": "not_performed",
              "self_test": "pending" if self_test else "not_requested"}
    if not report["python_supported"]:
        report["errors"].append("Python 3.10+ required for supported script execution")
    manifest = ROOT / "manifest.json"
    if manifest.is_file():
        try:
            document = json.loads(manifest.read_text(encoding="utf-8"))
            if not isinstance(document, dict) or document.get("algorithm") != "sha256":
                raise ValueError("Manifest must be an object using sha256")
            inventory = document.get("files")
            if not isinstance(inventory, dict) or not inventory:
                raise ValueError("Manifest files must be a nonempty object")
            required = {"SKILL.md", "baseline.json", "assets/ustc-format.json"}
            required.update("scripts/" + name + ".py" for name in ("doctor", "audit_project", "audit_text", "audit_docx", "audit_citations", "build_table", "simulate_survey"))
            packaged = {p.relative_to(ROOT).as_posix() for p in ROOT.rglob("*")
                        if p.is_file() and "__pycache__" not in p.parts and p.name != "manifest.json"
                        and (p.suffix.lower() in {".md", ".json", ".yaml", ".py", ".txt", ".pdf", ".docx"}
                             or p.name == ".gitattributes")}
            missing = (required | packaged) - inventory.keys()
            if missing:
                report["errors"].append("Manifest omits resources: " + ", ".join(sorted(missing)))
            for name, expected in inventory.items():
                relative = PurePosixPath(name)
                if relative.is_absolute() or ".." in relative.parts or "\\" in name or ":" in name or name != relative.as_posix() or name == "manifest.json":
                    report["errors"].append(f"Unsafe manifest path: {name}")
                    continue
                if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-f]{64}", expected):
                    report["errors"].append(f"Invalid SHA256 value: {name}")
                    continue
                path = (ROOT / name).resolve()
                if not path.is_relative_to(ROOT) or not path.is_file():
                    report["errors"].append(f"Missing or unsafe manifest path: {name}")
                elif hashlib.sha256(path.read_bytes()).hexdigest() != expected:
                    report["errors"].append(f"Checksum mismatch: {name}")
            report["checks"].append({"manifest_files": len(inventory)})
        except (OSError, ValueError, KeyError, TypeError) as exc:
            report["errors"].append(f"Invalid manifest: {exc}")
    else:
        report["errors"].append("manifest.json missing from baseline package")
    docx_available = importlib.util.find_spec("docx") is not None
    report["optional"]["python_docx"] = "available" if docx_available else "missing; audits still work"
    report["optional"]["renderer_candidates_on_path"] = {
        name: shutil.which(name) for name in ("soffice", "libreoffice", "winword")}
    report["optional"]["renderer_note"] = "PATH discovery only; native/API renderers may exist elsewhere. No renderer has been launched."
    if not self_test:
        return report
    if report["errors"]:
        report["self_test"] = "blocked_by_installation_errors"
        return report
    with tempfile.TemporaryDirectory(prefix="mba skill 自检 ") as temp:
        workspace = Path(temp)
        def run(name, args, expected=0):
            interpreter_flags = ["-S"] if sys.flags.no_site else []
            if sys.flags.optimize:
                interpreter_flags.append("-" + "O" * sys.flags.optimize)
            process = subprocess.run([sys.executable, *interpreter_flags, str(ROOT / "scripts" / name), *map(str, args)],
                                     cwd=workspace, capture_output=True, timeout=60,
                                     env=dict(os.environ, PYTHONIOENCODING="utf-8"))
            stdout = process.stdout.decode("utf-8")
            if process.returncode != expected:
                raise ValueError(f"{name} exit {process.returncode}: {stdout[:500]} {process.stderr.decode('utf-8', errors='replace')[:300]}")
            payload = json.loads(stdout)
            report["checks"].append({"script": name, "exit": process.returncode})
            return payload
        def verify(condition, message):
            # Assertions can disappear under Python -O/PYTHONOPTIMIZE.
            if not condition:
                raise ValueError(message)
        try:
            project = workspace / "项目.json"
            project.write_text(json.dumps({"schema_version": 2, "project": {
                "title": "SYNTHETIC self-test", "mode": "topic_to_thesis", "stage": "research_draft"}}, ensure_ascii=False), encoding="utf-8")
            verify(not run("audit_project.py", [project])["errors"], "Draft rejected")
            paragraph = "这是临时安装自检使用的模拟段落，内容用于核验重复扫描是否能够按实际段落运行，不代表任何企业事实或论文结论。" * 2
            prose = workspace / "测试.md"
            prose.write_text(paragraph + "\n\n" + paragraph, encoding="utf-8")
            result = run("audit_text.py", [prose])
            verify(any(x["kind"] == "exact" for x in result["matches"]), "Duplicate missed")
            citations = workspace / "引用自检.md"
            citations.write_text("模拟论述[2]。\n# 参考文献\n[1] SYNTHETIC test entry\n", encoding="utf-8")
            result = run("audit_citations.py", [citations, "--style", "numeric"], expected=1)
            verify(result.get("missing_labels") == [2], "Actual manuscript citation mismatch missed")
            survey = workspace / "模拟配置.json"
            survey.write_text(json.dumps({"data_origin": "SIMULATED", "purpose": "安装自检", "sample_size": 3, "seed": 7,
                "questions": [{"id": "Q1", "text": "自检选项", "kind": "single", "options": [{"value": "A", "label": "甲"}, {"value": "B", "label": "乙"}], "counts": [1, 2]}]}, ensure_ascii=False), encoding="utf-8")
            survey_out = workspace / "模拟输出"
            run("simulate_survey.py", [survey, "--out-dir", survey_out])
            stats = json.loads((survey_out / "simulated-summary.json").read_text(encoding="utf-8"))
            verify([o["count"] for o in stats["questions"][0]["overall"]["options"]] == [1, 2], "Simulation quotas failed")
            verify(stats["data_origin"] == "SIMULATED", "Simulation provenance missing")
            # Exercise full-thesis requirements even without optional docx libraries.
            namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            ET.register_namespace("w", namespace)
            tag = lambda name: "{" + namespace + "}" + name
            def thesis_fixture(path, english_text):
                document_xml = ET.Element(tag("document"))
                body_xml = ET.SubElement(document_xml, tag("body"))
                def paragraph(text, style=None):
                    node = ET.SubElement(body_xml, tag("p"))
                    if style:
                        ET.SubElement(ET.SubElement(node, tag("pPr")), tag("pStyle"), {tag("val"): style})
                    ET.SubElement(ET.SubElement(node, tag("r")), tag("t")).text = text
                    return node
                paragraph("摘要", "heading")
                paragraph("本模拟研究使用教学记录比较服务表现，分析责任接续并形成改进方案。")
                paragraph("关键词：服务；管理")
                paragraph("ABSTRACT", "heading")
                paragraph(english_text)
                paragraph("KEYWORDS: service; management")
                paragraph("目录", "heading")
                toc = paragraph("第1章 绪论\t1", "toc1")
                for kind, text in [("begin", None), ("instruction", ' TOC \\o "1-3" '), ("separate", None)]:
                    run_xml = ET.Element(tag("r"))
                    if kind == "instruction":
                        ET.SubElement(run_xml, tag("instrText")).text = text
                    else:
                        ET.SubElement(run_xml, tag("fldChar"), {tag("fldCharType"): kind})
                    toc.insert(len(toc) - 1, run_xml)
                ET.SubElement(ET.SubElement(toc, tag("r")), tag("fldChar"), {tag("fldCharType"): "end"})
                paragraph("第1章 绪论", "heading")
                paragraph("This synthetic self-test statement cites its test reference[1].")
                paragraph("参考文献", "heading")
                paragraph("[1] SYNTHETIC AUTHOR. Self-test reference[J]. Synthetic Journal, 2026, 1: 1.")
                styles_xml = ET.Element(tag("styles"))
                for sid, name, outline in [("heading", "Heading 1", "0"), ("toc1", "TOC 1", None)]:
                    style_xml = ET.SubElement(styles_xml, tag("style"), {tag("type"): "paragraph", tag("styleId"): sid})
                    ET.SubElement(style_xml, tag("name"), {tag("val"): name})
                    if outline is not None:
                        ET.SubElement(ET.SubElement(style_xml, tag("pPr")), tag("outlineLvl"), {tag("val"): outline})
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("word/document.xml", ET.tostring(document_xml, encoding="utf-8"))
                    archive.writestr("word/styles.xml", ET.tostring(styles_xml, encoding="utf-8"))
            complete = workspace / "完整性自检.docx"
            thesis_fixture(complete, "This synthetic teaching study compares service records and develops a management proposal.")
            result = run("audit_docx.py", [complete, "--full-thesis"])
            verify(all(item["status"] == "present" for item in result["required_content"].values()), "Required thesis content was not located")
            incomplete = workspace / "空摘要自检.docx"
            thesis_fixture(incomplete, "")
            result = run("audit_docx.py", [incomplete, "--full-thesis"], expected=1)
            verify(result["required_content"]["abstract_en"]["status"] == "empty", "Empty English abstract was not detected")
            if docx_available:
                from docx import Document
                spec = workspace / "表.json"
                spec.write_text(json.dumps({"number": "3.1", "title": "安装自检模拟表", "headers": ["类别", "值"],
                    "rows": [["模拟A", 1], ["模拟B", None]], "source": "临时生成的自检材料，非研究数据"}, ensure_ascii=False), encoding="utf-8")
                output = workspace / "表.docx"
                run("build_table.py", [spec, "--out", output])
                doc = Document(output)
                verify(len(doc.tables) == 1 and doc.tables[0].cell(2, 1).text == "—", "Editable table or missing value failed")
                verify(any("不表示零" in p.text for p in doc.paragraphs), "Missing value note omitted")
                result = run("audit_docx.py", [output, "--profile", ROOT / "assets/ustc-format.json"])
                verify(not result["errors"], "Generated table has structural errors")
                original = output.read_bytes()
                run("build_table.py", [spec, "--out", output], expected=2)
                verify(output.read_bytes() == original, "Existing output changed")
            else:
                # The standard-library audit can inspect the bundled DOCX without python-docx.
                run("audit_docx.py", [ROOT / "assets/ustc-thesis-template.docx"])
                report["checks"].append({"table_generation": "skipped: optional python-docx unavailable"})
        except (OSError, ValueError, ImportError, KeyError, TypeError, IndexError, subprocess.SubprocessError) as exc:
            report["errors"].append(str(exc))
    report["self_test"] = "failed" if report["errors"] else "passed"
    return report


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    try:
        result = inspect(args.self_test)
    except (OSError, ValueError, ImportError) as exc:
        result = {"errors": [str(exc)], "self_test": "failed", "visual_review": "not_performed"}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
