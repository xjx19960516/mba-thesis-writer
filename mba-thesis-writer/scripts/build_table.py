#!/usr/bin/env python3
"""Render supplied, verified data as editable DOCX tables. Requires python-docx."""
import argparse
import json
import math
from pathlib import Path
import re
import sys


def build(spec, profile, output, input_path=None, anchor=None, section_number=1):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    if not isinstance(spec, dict):
        raise ValueError("Table spec must be an object")
    for key in ["number", "title", "source"]:
        if not isinstance(spec.get(key), str) or not spec[key].strip():
            raise ValueError(f"{key} must contain meaningful text")
    if not re.fullmatch(r"(?:[1-9]\d*|[A-Z])\.[1-9]\d*", spec["number"]):
        raise ValueError("number must be chapter.index or appendix.index, such as 3.1 or A.1")
    if spec["title"].rstrip().endswith(tuple("。．.!！?？;；:：")):
        raise ValueError("Caption title must not end with punctuation")
    headers, rows = spec.get("headers"), spec.get("rows")
    if not isinstance(headers, list) or not headers or any(not isinstance(x, str) or not x.strip() for x in headers):
        raise ValueError("headers must be a nonempty string array")
    if not isinstance(rows, list) or not rows or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise ValueError("All data rows must have the header column count")
    for row in rows:
        for value in row:
            if value is not None and type(value) not in (str, int, float):
                raise ValueError("Cells accept strings, finite numbers or null")
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError("Non-finite numeric cell")
    for key in ["continuation", "new_page"]:
        if key in spec and type(spec[key]) is not bool:
            raise ValueError(f"{key} must be boolean")
    for key in ["note"]:
        if key in spec and not isinstance(spec[key], str):
            raise ValueError(f"{key} must be a string")
    if output.exists():
        raise ValueError("Output already exists; choose a new output path")
    doc = Document(input_path) if input_path else Document()
    if not 1 <= section_number <= len(doc.sections):
        raise ValueError("Invalid 1-based section number")
    anchor_p = None
    if input_path:
        matches = [p for p in doc.paragraphs if p.text == anchor]
        if not anchor or len(matches) != 1:
            raise ValueError("Exactly one complete top-level anchor paragraph required")
        anchor_p = matches[0]
        if anchor_p._p.find("w:pPr/w:sectPr", anchor_p._p.nsmap) is not None:
            raise ValueError("Anchor contains a section break; use a plain paragraph")
        # A section's properties are located at its end, after the anchor.
        preceding = 0
        for el in doc.element.body:
            if el is anchor_p._p:
                break
            if el.find(".//w:sectPr", el.nsmap) is not None:
                preceding += 1
        if preceding + 1 != section_number:
            raise ValueError("--section does not match anchor section")
    section = doc.sections[section_number - 1]
    if not input_path:
        pg = profile["page"]
        for attr, key in [("page_width", "width_cm"), ("page_height", "height_cm"), ("top_margin", "top_cm"), ("bottom_margin", "bottom_cm"), ("left_margin", "left_cm"), ("right_margin", "right_cm"), ("gutter", "gutter_cm"), ("header_distance", "header_cm"), ("footer_distance", "footer_cm")]:
            setattr(section, attr, Cm(pg[key]))
        if pg.get("mirror"):
            doc.settings.element.append(OxmlElement("w:mirrorMargins"))
    width = (section.page_width - section.left_margin - section.right_margin - section.gutter) / 360000
    widths = spec.get("column_widths_cm", [width / len(headers)] * len(headers))
    if not isinstance(widths, list) or len(widths) != len(headers) or any(type(x) not in (int, float) or not math.isfinite(x) or x <= 0 for x in widths) or sum(widths) > width + .01:
        raise ValueError("Column widths must be positive and fit the actual section text width")
    aligns = spec.get("alignments", ["left"] * len(headers))
    if not isinstance(aligns, list) or len(aligns) != len(headers) or any(x not in ("left", "center", "right") for x in aligns):
        raise ValueError("alignments must match columns and use left/center/right")
    border_style = spec.get("border_style", "three_line")
    if border_style not in ("three_line", "grid", "none"):
        raise ValueError("Unknown border_style")
    cfg = profile["table"]
    padding_cm = cfg.get("cell_padding_cm", 0.15)
    if type(padding_cm) not in (int, float) or not math.isfinite(padding_cm) or padding_cm < 0 or min(widths) <= padding_cm * 2:
        raise ValueError("Cell padding must be finite, nonnegative and leave usable column width")
    new_nodes = []
    # Keep host styles unchanged; isolate inserted content from a custom Normal.
    style_name = "MBA Table Text"
    suffix = 1
    while style_name in doc.styles:
        suffix += 1
        style_name = f"MBA Table Text {suffix}"
    text_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    text_style.base_style = None
    text_style.font.name = cfg["latin"]
    text_style.font.size = Pt(cfg["size_pt"])
    text_style.font.italic = False
    text_style.font.underline = False

    def para(text="", before=0, after=0, bold=False, center=False):
        p = doc.add_paragraph(style=text_style)
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = pf.right_indent = Pt(0)
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.keep_with_next = False
        pf.page_break_before = False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        font(p.add_run(text), bold)
        new_nodes.append(p._p)
        return p

    def font(run, bold=False):
        run.font.name = cfg["latin"]
        run.font.size = Pt(cfg["size_pt"])
        run.font.bold = bold
        run.font.italic = False
        run.font.underline = False
        run.font.color.rgb = RGBColor(0, 0, 0)
        rp = run._element.get_or_add_rPr()
        rfonts = rp.find(qn("w:rFonts"))
        for key, value in [("eastAsia", cfg["east_asia"]), ("ascii", cfg["latin"]), ("hAnsi", cfg["latin"])]:
            rfonts.set(qn("w:" + key), value)

    gap = para()
    gap.paragraph_format.line_spacing = Pt(cfg["surrounding_blank_pt"])
    gap.paragraph_format.keep_with_next = True
    gap.paragraph_format.page_break_before = spec.get("new_page", False)
    caption = para(("续表" if spec.get("continuation") else "表") + spec["number"] + "　" + spec["title"], before=cfg["caption_before_pt"], after=cfg["caption_after_pt"], bold=True, center=True)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table_style_name = "MBA Data Table"
    suffix = 1
    while table_style_name in doc.styles:
        suffix += 1
        table_style_name = f"MBA Data Table {suffix}"
    table_style = doc.styles.add_style(table_style_name, WD_STYLE_TYPE.TABLE)
    table_style.base_style = None
    table.style = table_style
    new_nodes.append(table._tbl)
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    for col, value in zip(table.columns, widths):
        col.width = Cm(value)
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = "—" if value is None else str(value)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        line = OxmlElement("w:" + edge)
        enabled = border_style == "grid" or (border_style == "three_line" and edge in ("top", "bottom"))
        line.set(qn("w:val"), "single" if enabled else "nil")
        line.set(qn("w:sz"), "8")
        line.set(qn("w:color"), "000000")
        borders.append(line)
    table._tbl.tblPr.append(borders)
    # Isolated table styles have no inherited cell margins in some renderers.
    margins = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        margin = OxmlElement("w:" + edge)
        margin.set(qn("w:w"), str(round(padding_cm * 1440 / 2.54)))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.append(margins)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            trpr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            trpr.append(header)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.style = text_style
                pf = p.paragraph_format
                pf.first_line_indent = pf.left_indent = pf.right_indent = Pt(0)
                pf.space_before, pf.space_after = Pt(cfg["before_pt"]), Pt(cfg["after_pt"])
                pf.line_spacing = 1
                pf.keep_with_next = row_index == 0
                pf.page_break_before = False
                p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}["center" if row_index == 0 else aligns[col_index]]
                for run in p.runs:
                    font(run, row_index == 0)
            if row_index == 0 and border_style == "three_line":
                cell_borders = OxmlElement("w:tcBorders")
                line = OxmlElement("w:bottom")
                line.set(qn("w:val"), "single")
                line.set(qn("w:sz"), "6")
                cell_borders.append(line)
                cell._tc.get_or_add_tcPr().append(cell_borders)
    para("资料来源：" + spec["source"], before=cfg["note_before_pt"], after=cfg["note_after_pt"])
    if spec.get("note") or any(value is None for row in rows for value in row):
        note = spec.get("note", "")
        if any(value is None for row in rows for value in row):
            note += ("；" if note else "") + "—表示缺失或未披露，不表示零"
        para("注：" + note, before=cfg["note_before_pt"], after=cfg["note_after_pt"])
    gap = para()
    gap.paragraph_format.line_spacing = Pt(cfg["surrounding_blank_pt"])
    if anchor_p is not None:
        for node in new_nodes:
            anchor_p._p.addprevious(node)
        anchor_p._p.getparent().remove(anchor_p._p)
    doc.save(output)
    return {"output": str(output.resolve()), "rows": len(rows), "columns": len(headers), "visual_review": "pending", "continuation_layout": "render then verify split and caption", "border_authority": "user setting or skill default, not mandatory USTC rule"}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--input", type=Path)
    parser.add_argument("--anchor")
    parser.add_argument("--section", type=int, default=1)
    parser.add_argument("--profile", type=Path, default=Path(__file__).resolve().parents[1] / "assets" / "ustc-format.json")
    args = parser.parse_args()
    if args.input and args.input.resolve() == args.out.resolve():
        parser.error("Input must remain unchanged; select a new output")
    try:
        result = build(json.loads(args.spec.read_text(encoding="utf-8-sig")), json.loads(args.profile.read_text(encoding="utf-8-sig")), args.out, args.input, args.anchor, args.section)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ImportError, OSError, ValueError, KeyError, TypeError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    sys.exit(main())
